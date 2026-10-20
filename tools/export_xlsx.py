"""Write the site's workbooks (tools/cache/*.json) as .xlsx – same cells, dates as serials with a date format.
Used to run the reference formulas through other spreadsheet engines (and to feed GridCraft)."""
import json, os, sys, glob, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font
HERE = os.path.dirname(os.path.abspath(__file__))
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, 'cache', 'xlsx')
os.makedirs(OUT, exist_ok=True)
info = {}
for f in sorted(glob.glob(os.path.join(HERE, 'cache', '*.json'))):
    data = json.load(open(f))
    wb = Workbook(); wb.remove(wb.active)
    for s in data['sheets']:
        ws = wb.create_sheet(s['name'][:31])
        dcols = {int(k): v for k, v in s.get('dateCols', {}).items()}
        for i, row in enumerate(s['rows']):
            for j, v in enumerate(row):
                if v is None: continue
                c = ws.cell(row=i + 1, column=j + 1, value=v)
                if i == 0: c.font = Font(bold=True)
                elif j in dcols and isinstance(v, (int, float)): c.number_format = dcols[j]
                if j in (s.get('textCols') or []) and i > 0: c.data_type = 's'; c.value = str(v)
    info[data['id']] = [{'name': s['name'], 'headers': s['rows'][0], 'rows': len(s['rows']) - 1, 'example': s['rows'][1] if len(s['rows']) > 1 else None} for s in data['sheets']]
    wb.create_sheet('Work')
    p = os.path.join(OUT, data['id'] + '.xlsx'); wb.save(p); print(p, os.path.getsize(p) // 1024, 'KB')
json.dump(info, open(os.path.join(HERE, '..', 'src', 'data', 'wbinfo.json'), 'w'), default=str)
